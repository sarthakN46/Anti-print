import sys
import json
import os
import traceback
import subprocess
import shutil

def get_pdf_page_count(pdf_path):
    from pypdf import PdfReader
    reader = PdfReader(pdf_path)
    return len(reader.pages)

def convert_to_pdf_temp(input_path):
    """
    Converts file to a temporary PDF using LibreOffice (Headless).
    Returns path to generated PDF or None if failed.
    """
    try:
        # Check for libreoffice
        cmd = 'libreoffice'
        if shutil.which('libreoffice') is None:
            if shutil.which('soffice') is not None:
                cmd = 'soffice'
            else:
                return None

        out_dir = os.path.dirname(input_path)
        
        args = [
            cmd,
            '--headless',
            '--convert-to',
            'pdf',
            '--outdir',
            out_dir,
            input_path
        ]
        
        # Run conversion
        subprocess.run(args, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=20)
        
        # Expected output file
        base_name = os.path.splitext(os.path.basename(input_path))[0]
        pdf_path = os.path.join(out_dir, base_name + '.pdf')
        
        if os.path.exists(pdf_path):
            return pdf_path
        return None
        
    except Exception:
        return None

def analyze_file(file_path):
    try:
        file_path = os.path.abspath(file_path)
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()
        page_count = 1
        file_type = 'unknown'

        # --- PDF ---
        if ext == '.pdf':
            file_type = 'pdf'
            page_count = get_pdf_page_count(file_path)
        
        # --- OFFICE DOCS (DOCX, DOC, PPTX, PPT, ODT) ---
        elif ext in ['.docx', '.doc', '.pptx', '.ppt', '.odt']:
            file_type = ext.replace('.', '')
            # Try accurate count via PDF conversion (LibreOffice)
            pdf_temp = convert_to_pdf_temp(file_path)
            
            if pdf_temp:
                try:
                    page_count = get_pdf_page_count(pdf_temp)
                except:
                    page_count = 1 # Fallback
                finally:
                    if os.path.exists(pdf_temp):
                        os.remove(pdf_temp)
            else:
                # Fallback: Approximate or Default
                if ext == '.docx':
                    try:
                        from docx import Document
                        doc = Document(file_path)
                        if doc.core_properties.page_count:
                            page_count = doc.core_properties.page_count
                    except:
                        pass
                elif ext == '.pptx':
                    try:
                        from pptx import Presentation
                        prs = Presentation(file_path)
                        page_count = len(prs.slides)
                    except:
                        pass
        
        # --- EXCEL / CSV ---
        elif ext in ['.xlsx', '.xls', '.csv']:
            file_type = 'excel'
            # Excel page count is tricky (depends on print area/scaling).
            # Default to 1 is standard behavior unless we render it.
            # We COULD render it to PDF, but it might be huge.
            # Let's keep 1 for Excel for safety/speed.
            page_count = 1
            
        # --- IMAGES ---
        elif ext in ['.png', '.jpg', '.jpeg']:
            file_type = 'image'
            page_count = 1

        print(json.dumps({"pageCount": page_count, "type": file_type}))

    except Exception as e:
        sys.stderr.write(f"Error analyzing file: {str(e)}\n")
        sys.stderr.write(traceback.format_exc())
        print(json.dumps({"pageCount": 1, "type": "unknown", "error": str(e)}))
        sys.exit(1)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(json.dumps({"error": "No file provided"}))
        sys.exit(1)
    
    analyze_file(sys.argv[1])